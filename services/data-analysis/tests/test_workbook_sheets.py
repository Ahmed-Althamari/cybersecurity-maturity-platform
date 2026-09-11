import io

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

from analysis.workbook_sheets import extract_workbook_sheets


def _workbook_bytes(wb: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _mixed_workbook() -> bytes:
    """One data table sheet, one dashboard sheet with a bar chart of two series."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RawData"
    ws.append(["A", "B"])
    ws.append([1, 2])
    ws.append([3, 4])

    dashboard = wb.create_sheet("Dashboard")
    dashboard.append(["Month", "Sales", "Costs"])
    for month, sales, costs in [("Jan", 100, 40), ("Feb", 150, 60), ("Mar", 130, 50)]:
        dashboard.append([month, sales, costs])

    chart = BarChart()
    chart.title = "Sales vs Costs"
    cats = Reference(dashboard, min_col=1, min_row=2, max_row=4)
    vals = Reference(dashboard, min_col=2, max_col=3, min_row=1, max_row=4)
    chart.add_data(vals, titles_from_data=True)
    chart.set_categories(cats)
    dashboard.add_chart(chart, "E2")

    return _workbook_bytes(wb)


def test_classifies_table_and_dashboard_sheets():
    result = extract_workbook_sheets("workbook.xlsx", _mixed_workbook())

    by_name = {s.name: s for s in result.sheets}
    assert by_name["RawData"].type == "TABLE"
    assert by_name["RawData"].row_count == 2  # header excluded
    assert by_name["Dashboard"].type == "DASHBOARD"


def test_extracts_real_chart_data_not_just_an_image():
    result = extract_workbook_sheets("workbook.xlsx", _mixed_workbook())

    assert len(result.extracted_charts) == 1
    chart = result.extracted_charts[0]
    assert chart.sheet_name == "Dashboard"
    assert chart.title == "Sales vs Costs"
    assert chart.chart_type == "BarChart"
    assert chart.categories == ["Jan", "Feb", "Mar"]
    series_by_name = {s.name: s.values for s in chart.series}
    assert series_by_name == {"Sales": [100.0, 150.0, 130.0], "Costs": [40.0, 60.0, 50.0]}
    assert len(chart.image_base64) > 100


def test_pie_and_line_charts_extract_and_render():
    wb = openpyxl.Workbook()
    pie_ws = wb.active
    pie_ws.title = "Pie"
    pie_ws.append(["Category", "Share"])
    for cat, share in [("A", 30), ("B", 50), ("C", 20)]:
        pie_ws.append([cat, share])
    pie = PieChart()
    pie.title = "Share breakdown"
    pie.add_data(Reference(pie_ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    pie.set_categories(Reference(pie_ws, min_col=1, min_row=2, max_row=4))
    pie_ws.add_chart(pie, "D2")

    line_ws = wb.create_sheet("Trend")
    line_ws.append(["Month", "Value"])
    for i, month in enumerate(["Jan", "Feb", "Mar", "Apr"]):
        line_ws.append([month, i * 10])
    line = LineChart()
    line.title = "Trend over time"
    line.add_data(Reference(line_ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    line.set_categories(Reference(line_ws, min_col=1, min_row=2, max_row=5))
    line_ws.add_chart(line, "D2")

    result = extract_workbook_sheets("workbook.xlsx", _workbook_bytes(wb))

    charts_by_type = {c.chart_type: c for c in result.extracted_charts}
    assert charts_by_type["PieChart"].series[0].values == [30.0, 50.0, 20.0]
    assert charts_by_type["LineChart"].series[0].values == [0.0, 10.0, 20.0, 30.0]


def test_csv_is_a_single_table_sheet_with_no_charts():
    df = pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]})
    csv_bytes = df.to_csv(index=False).encode("utf-8")

    result = extract_workbook_sheets("data.csv", csv_bytes)

    assert len(result.sheets) == 1
    assert result.sheets[0] == result.sheets[0]  # sanity
    assert result.sheets[0].type == "TABLE"
    assert result.sheets[0].row_count == 3
    assert result.sheets[0].column_count == 2
    assert result.extracted_charts == []


def test_unsupported_extension_raises_value_error():
    try:
        extract_workbook_sheets("data.txt", b"not a spreadsheet")
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_workbook_with_no_charts_has_only_table_sheets():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plain"
    ws.append(["x", "y"])
    ws.append([1, 2])

    result = extract_workbook_sheets("workbook.xlsx", _workbook_bytes(wb))

    assert all(s.type == "TABLE" for s in result.sheets)
    assert result.extracted_charts == []
