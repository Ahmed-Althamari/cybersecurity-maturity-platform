import io
import json

import numpy as np
import openpyxl
import pandas as pd
from fastapi.testclient import TestClient
from openpyxl.chart import BarChart, Reference

import main
from analysis.ai_mode import AiAnalysisResult, LlmNotConfiguredError


def _sample_csv_bytes() -> bytes:
    rng = np.random.default_rng(0)
    df = pd.DataFrame(
        {
            "category": rng.choice(["A", "B", "C"], 40),
            "maturity": rng.random(40) * 5,
        }
    )
    buf = io.BytesIO()
    df.to_csv(buf, index=False)
    return buf.getvalue()


def _multi_sheet_workbook_bytes() -> bytes:
    """A "SheetA" data table and a "SheetB" data table, plus a "Dashboard" sheet with a chart."""
    wb = openpyxl.Workbook()
    sheet_a = wb.active
    sheet_a.title = "SheetA"
    sheet_a.append(["x", "y"])
    sheet_a.append([1, 10])
    sheet_a.append([2, 20])

    sheet_b = wb.create_sheet("SheetB")
    sheet_b.append(["x", "y"])
    sheet_b.append([100, 1000])

    dashboard = wb.create_sheet("Dashboard")
    dashboard.append(["Month", "Value"])
    dashboard.append(["Jan", 5])
    dashboard.append(["Feb", 9])
    chart = BarChart()
    chart.title = "Monthly value"
    chart.add_data(Reference(dashboard, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(dashboard, min_col=1, min_row=2, max_row=3))
    dashboard.add_chart(chart, "D2")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


client = TestClient(main.app)


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_rejects_unknown_mode():
    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "bogus"},
    )
    assert response.status_code == 400


def test_rejects_unsupported_file_type():
    response = client.post(
        "/analyze",
        files={"file": ("data.txt", b"not a spreadsheet", "text/plain")},
        data={"mode": "local"},
    )
    assert response.status_code == 400


def test_local_mode_end_to_end():
    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "local"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["mode"] == "local"
    assert body["rowCount"] == 40
    assert body["columnCount"] == 2
    assert len(body["charts"]) > 0
    assert body["error"] is None


def test_ai_mode_returns_503_when_llm_not_configured(monkeypatch):
    def raise_not_configured(df, question, providers_override=None):
        raise LlmNotConfiguredError("no provider configured")

    monkeypatch.setattr(main, "run_ai_analysis", raise_not_configured)

    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "ai"},
    )
    assert response.status_code == 503


def test_ai_mode_success_uses_configured_llm(monkeypatch):
    def fake_run_ai_analysis(df, question, providers_override=None):
        return AiAnalysisResult(answer="the data looks fine", table=None, chart_base64=None, error=None)

    monkeypatch.setattr(main, "run_ai_analysis", fake_run_ai_analysis)

    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "ai", "question": "Summarize this"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["answer"] == "the data looks fine"
    assert body["charts"] == []


def test_ai_mode_forwards_llm_providers_form_field(monkeypatch):
    received = {}

    def fake_run_ai_analysis(df, question, providers_override=None):
        received["providers_override"] = providers_override
        return AiAnalysisResult(answer="ok", table=None, chart_base64=None, error=None)

    monkeypatch.setattr(main, "run_ai_analysis", fake_run_ai_analysis)

    providers = [{"format": "anthropic", "apiKey": "sk-ant-test", "model": "claude-opus-5", "baseUrl": None}]
    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "ai", "llm_providers": json.dumps(providers)},
    )
    assert response.status_code == 200
    assert received["providers_override"] == providers


def test_ai_mode_ignores_malformed_llm_providers_field(monkeypatch):
    received = {}

    def fake_run_ai_analysis(df, question, providers_override=None):
        received["providers_override"] = providers_override
        return AiAnalysisResult(answer="ok", table=None, chart_base64=None, error=None)

    monkeypatch.setattr(main, "run_ai_analysis", fake_run_ai_analysis)

    response = client.post(
        "/analyze",
        files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")},
        data={"mode": "ai", "llm_providers": "not valid json"},
    )
    assert response.status_code == 200
    assert received["providers_override"] is None


def test_sheets_classifies_and_extracts_charts_from_a_multi_sheet_workbook():
    response = client.post(
        "/sheets",
        files={"file": ("workbook.xlsx", _multi_sheet_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()

    by_name = {s["name"]: s for s in body["sheets"]}
    assert by_name["SheetA"]["type"] == "TABLE"
    assert by_name["SheetB"]["type"] == "TABLE"
    assert by_name["Dashboard"]["type"] == "DASHBOARD"

    assert len(body["extractedCharts"]) == 1
    chart = body["extractedCharts"][0]
    assert chart["sheetName"] == "Dashboard"
    assert chart["title"] == "Monthly value"
    assert chart["categories"] == ["Jan", "Feb"]
    assert chart["series"] == [{"name": "Value", "values": [5.0, 9.0]}]
    assert len(chart["imageBase64"]) > 100


def test_sheets_treats_a_csv_as_one_table_sheet():
    response = client.post("/sheets", files={"file": ("data.csv", _sample_csv_bytes(), "text/csv")})
    assert response.status_code == 200
    body = response.json()
    assert len(body["sheets"]) == 1
    assert body["sheets"][0]["type"] == "TABLE"
    assert body["extractedCharts"] == []


def test_sheets_rejects_unsupported_file_type():
    response = client.post("/sheets", files={"file": ("data.txt", b"not a spreadsheet", "text/plain")})
    assert response.status_code == 400


def test_analyze_reads_the_requested_sheet_not_the_first_one():
    response = client.post(
        "/analyze",
        files={"file": ("workbook.xlsx", _multi_sheet_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"mode": "local", "sheet_name": "SheetB"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["rowCount"] == 1  # SheetB has one data row; SheetA (the default) has two


def test_analyze_rejects_a_sheet_name_that_does_not_exist():
    response = client.post(
        "/analyze",
        files={"file": ("workbook.xlsx", _multi_sheet_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"mode": "local", "sheet_name": "NoSuchSheet"},
    )
    assert response.status_code == 400
