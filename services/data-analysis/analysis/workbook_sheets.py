"""Sheet classification + embedded-chart extraction for uploaded workbooks.

A real spreadsheet upload is rarely one clean data table: it commonly has a mix of raw "data"
sheets and one or more "dashboard" sheets the user (or whatever tool produced the file) already
built, containing native Excel chart objects and/or pivot tables rather than plain rows. AutoViz
and PandasAI (see local_mode.py/ai_mode.py) only know how to analyze a single flat DataFrame, so
neither can make sense of a dashboard sheet — this module lets the caller tell the two apart
*before* choosing which sheet to hand to the existing analysis modes, and pulls the dashboard
sheets' own charts out directly instead.

Classification is purely structural (no LLM call, no cost, works with no AI provider configured
at all): a sheet is "DASHBOARD" if it contains a native chart object and/or a pivot table,
otherwise "TABLE". This is a deliberate scope choice — see docs/security-architecture.md's
pattern of preferring a cheap deterministic signal over an AI call wherever one exists.

Chart extraction reads the *real* values a chart plots, not just its picture: a chart's series
reference (e.g. `'Sheet1'!B2:B10`) carries a cached snapshot of the values Excel last rendered
(`numCache`/`strCache`); when present we trust it (it's exactly what the user saw), and when it's
missing or empty (e.g. a workbook written by a tool — including this module's own test fixtures —
that never triggered a recalculation) we fall back to reading the referenced cell range directly.
Either way, the caller gets real category/series numbers back, not just an image — which is what
lets a chart pinned to the dashboard later support a genuine "click through to the data" view,
unlike AutoViz/PandasAI's charts, which are images with no recoverable structure.
"""

from __future__ import annotations

import base64
import io
import logging
import re
from dataclasses import dataclass, field

import matplotlib

matplotlib.use("Agg")  # headless: no display available in a server process

import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger("data-analysis.workbook_sheets")

# A chart's data-source formula looks like "'Sheet Name'!$B$2:$B$10" or "Sheet1!B2:B10".
_REF_PATTERN = re.compile(r"^(?:'([^']+)'|([^!']+))!(.+)$")

# Chart types this module knows how to re-render; anything else falls back to a grouped bar chart
# rather than failing extraction outright — the real category/series data is still returned either
# way, only the re-rendered picture's shape is a fallback.
_PIE_TYPES = {"PieChart", "DoughnutChart"}
_LINE_TYPES = {"LineChart", "AreaChart"}


@dataclass
class SheetSummary:
    name: str
    row_count: int
    column_count: int
    type: str  # "TABLE" | "DASHBOARD"


@dataclass
class ExtractedSeries:
    name: str
    values: list[float]


@dataclass
class ExtractedChart:
    sheet_name: str
    title: str
    chart_type: str
    image_base64: str
    categories: list[str] = field(default_factory=list)
    series: list[ExtractedSeries] = field(default_factory=list)


@dataclass
class WorkbookSheetsResult:
    sheets: list[SheetSummary]
    extracted_charts: list[ExtractedChart]


def _parse_ref(formula: str | None) -> tuple[str, str] | None:
    if not formula:
        return None
    match = _REF_PATTERN.match(formula)
    if not match:
        return None
    return (match.group(1) or match.group(2), match.group(3))


def _read_ref_values(workbook: openpyxl.Workbook, formula: str | None) -> list:
    """Fallback path: read the actual cell values a chart's data-source formula points at,
    for when the chart carries no cached snapshot (or an empty one) to trust instead."""
    parsed = _parse_ref(formula)
    if parsed is None:
        return []
    sheet_name, cell_range = parsed
    if sheet_name not in workbook.sheetnames:
        return []
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError:
        return []
    values: list = []
    for row in workbook[sheet_name].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        values.extend(cell.value for cell in row)
    return values


def _cache_values(cache) -> list | None:
    """Returns None (not an empty list) when there's no usable cache, so the caller can tell
    "no cache, fall back to cell reads" apart from "cache says this series is genuinely empty"."""
    if cache is None or not cache.pt:
        return None
    size = cache.ptCount or (max(pt.idx for pt in cache.pt) + 1)
    values: list = [None] * size
    for pt in cache.pt:
        if 0 <= pt.idx < size:
            values[pt.idx] = pt.v
    return values


def _chart_title(chart, index: int) -> str:
    title = getattr(chart, "title", None)
    tx = getattr(title, "tx", None) if title is not None else None
    if tx is not None and tx.rich is not None:
        texts = [run.t for paragraph in tx.rich.p for run in (paragraph.r or []) if run.t]
        joined = "".join(texts).strip()
        if joined:
            return joined
    if tx is not None and tx.strRef is not None:
        cached = _cache_values(tx.strRef.strCache)
        if cached and cached[0]:
            return str(cached[0])
    return f"Chart {index + 1}"


def _series_name(workbook: openpyxl.Workbook, series, index: int) -> str:
    tx = series.tx
    if tx is None:
        return f"Series {index + 1}"
    if tx.strRef is not None:
        cached = _cache_values(tx.strRef.strCache)
        if cached and cached[0] is not None:
            return str(cached[0])
        resolved = _read_ref_values(workbook, tx.strRef.f)
        if resolved and resolved[0] is not None:
            return str(resolved[0])
    if getattr(tx, "v", None) is not None:
        return str(tx.v)
    return f"Series {index + 1}"


def _categories(workbook: openpyxl.Workbook, cat) -> list[str]:
    if cat is None:
        return []
    ref = cat.strRef or cat.numRef
    if ref is None:
        return []
    cache = getattr(ref, "strCache", None) or getattr(ref, "numCache", None)
    cached = _cache_values(cache)
    if cached is not None:
        return ["" if v is None else str(v) for v in cached]
    return ["" if v is None else str(v) for v in _read_ref_values(workbook, ref.f)]


def _series_values(workbook: openpyxl.Workbook, val) -> list[float]:
    if val is None or val.numRef is None:
        return []
    cached = _cache_values(val.numRef.numCache)
    source = cached if cached is not None else _read_ref_values(workbook, val.numRef.f)
    return [float(v) if isinstance(v, (int, float)) else 0.0 for v in source]


def _render_chart_png(title: str, chart_type: str, categories: list[str], series: list[ExtractedSeries]) -> str:
    fig, ax = plt.subplots(figsize=(6, 4))
    try:
        if chart_type in _PIE_TYPES and series:
            ax.pie(series[0].values, labels=categories or None, autopct="%1.0f%%")
        elif chart_type in _LINE_TYPES:
            for s in series:
                ax.plot(categories or range(len(s.values)), s.values, marker="o", label=s.name)
            if len(series) > 1:
                ax.legend()
        else:
            width = 0.8 / max(len(series), 1)
            x_positions = range(len(categories) or max((len(s.values) for s in series), default=0))
            for i, s in enumerate(series):
                offsets = [x + i * width for x in x_positions]
                ax.bar(offsets, s.values, width=width, label=s.name)
            if categories:
                ax.set_xticks([x + width * (len(series) - 1) / 2 for x in x_positions])
                ax.set_xticklabels(categories, rotation=45, ha="right")
            if len(series) > 1:
                ax.legend()
        ax.set_title(title)
        fig.tight_layout()
        buffer = io.BytesIO()
        fig.savefig(buffer, format="png")
        return base64.b64encode(buffer.getvalue()).decode("utf-8")
    finally:
        plt.close(fig)


def _extract_chart(workbook: openpyxl.Workbook, sheet_name: str, chart, index: int) -> ExtractedChart | None:
    title = _chart_title(chart, index)
    chart_type = type(chart).__name__
    series_list: list[ExtractedSeries] = []
    categories: list[str] = []
    for series_index, series in enumerate(chart.series):
        values = _series_values(workbook, series.val)
        if not values:
            continue
        series_list.append(ExtractedSeries(name=_series_name(workbook, series, series_index), values=values))
        if not categories:
            categories = _categories(workbook, series.cat)

    if not series_list:
        # Nothing plottable came back (e.g. a chart referencing a range that no longer exists) —
        # skip it rather than render a blank picture.
        return None

    image_base64 = _render_chart_png(title, chart_type, categories, series_list)
    return ExtractedChart(
        sheet_name=sheet_name,
        title=title,
        chart_type=chart_type,
        image_base64=image_base64,
        categories=categories,
        series=series_list,
    )


def _classify_sheet(ws: Worksheet) -> bool:
    """True if this sheet is a "dashboard" sheet — one built around a chart or pivot table
    rather than plain tabular data."""
    has_charts = bool(ws._charts)
    has_pivots = bool(getattr(ws, "_pivots", None))
    return has_charts or has_pivots


def extract_workbook_sheets(filename: str, content: bytes) -> WorkbookSheetsResult:
    lower = filename.lower()

    if lower.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content))
        return WorkbookSheetsResult(
            sheets=[SheetSummary(name="Sheet1", row_count=len(df), column_count=len(df.columns), type="TABLE")],
            extracted_charts=[],
        )

    if not (lower.endswith(".xlsx") or lower.endswith(".xls")):
        raise ValueError(f"Unsupported file type: {filename} (expected .csv, .xlsx or .xls)")

    if lower.endswith(".xls"):
        # openpyxl only reads the modern .xlsx format — legacy .xls has no chart/pivot
        # introspection available here, so it's treated as a single opaque table sheet (matching
        # /analyze's existing pd.read_excel behavior for this format).
        df = pd.read_excel(io.BytesIO(content))
        return WorkbookSheetsResult(
            sheets=[SheetSummary(name="Sheet1", row_count=len(df), column_count=len(df.columns), type="TABLE")],
            extracted_charts=[],
        )

    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets: list[SheetSummary] = []
    extracted_charts: list[ExtractedChart] = []

    for ws in workbook.worksheets:
        is_dashboard = _classify_sheet(ws)
        sheets.append(
            SheetSummary(
                name=ws.title,
                row_count=max(0, (ws.max_row or 0) - 1),
                column_count=ws.max_column or 0,
                type="DASHBOARD" if is_dashboard else "TABLE",
            )
        )
        for index, chart in enumerate(ws._charts):
            try:
                extracted = _extract_chart(workbook, ws.title, chart, index)
            except Exception:
                logger.exception("Failed to extract chart %d on sheet %r — skipping it", index, ws.title)
                continue
            if extracted is not None:
                extracted_charts.append(extracted)

    return WorkbookSheetsResult(sheets=sheets, extracted_charts=extracted_charts)
